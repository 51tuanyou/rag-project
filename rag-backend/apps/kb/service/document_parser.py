"""
Document parsing service for handling various file formats
"""
import os
import tempfile
from html import escape
from typing import List, Dict, Any
import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
import PyPDF2
import pandas as pd


class DocumentParser:
    def __init__(self):
        self.supported_formats = ['.docx', '.pdf', '.txt', '.csv']
    
    def parse_document(self, file_path: str) -> str:
        """Parse document and extract text content"""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_ext = os.path.splitext(file_path)[1].lower()
        
        if file_ext == '.docx':
            return self._parse_docx(file_path)
        elif file_ext == '.pdf':
            return self._parse_pdf(file_path)
        elif file_ext == '.txt':
            return self._parse_txt(file_path)
        elif file_ext == '.csv':
            return self._parse_csv(file_path)
        else:
            raise ValueError(f"Unsupported file format: {file_ext}")
    
    def _parse_docx(self, file_path: str) -> str:
        """Parse DOCX file"""
        try:
            doc = docx.Document(file_path)
            text_content = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text.strip())
            
            return '\n\n'.join(text_content)
        except Exception as e:
            raise Exception(f"Error parsing DOCX file: {str(e)}")
    
    def _parse_pdf(self, file_path: str) -> str:
        """Parse PDF file"""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                text_content = []
                
                print(f"PDF has {len(pdf_reader.pages)} pages")
                
                for i, page in enumerate(pdf_reader.pages):
                    try:
                        page_text = page.extract_text()
                        if page_text and page_text.strip():
                            text_content.append(page_text.strip())
                            print(f"Page {i+1}: {len(page_text)} characters")
                        else:
                            print(f"Page {i+1}: No text content")
                    except Exception as page_error:
                        print(f"Error extracting text from page {i+1}: {str(page_error)}")
                        continue
                
                if not text_content:
                    raise Exception("No text content found in PDF")
                
                result = '\n\n'.join(text_content)
                print(f"Total extracted text: {len(result)} characters")
                return result
        except Exception as e:
            print(f"PDF parsing error: {str(e)}")
            raise Exception(f"Error parsing PDF file: {str(e)}")
    
    def _parse_txt(self, file_path: str) -> str:
        """Parse TXT file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except Exception as e:
            raise Exception(f"Error parsing TXT file: {str(e)}")
    
    def _parse_csv(self, file_path: str) -> str:
        """Parse CSV file"""
        try:
            df = pd.read_csv(file_path)
            # Convert DataFrame to text format
            text_content = []
            for index, row in df.iterrows():
                row_text = ' | '.join([str(value) for value in row.values if pd.notna(value)])
                text_content.append(row_text)
            
            return '\n'.join(text_content)
        except Exception as e:
            raise Exception(f"Error parsing CSV file: {str(e)}")

    def render_browser_html(self, file_path: str, title: str) -> str | None:
        """HTML preview for formats browsers cannot display natively."""
        file_ext = os.path.splitext(file_path)[1].lower()
        if file_ext == '.docx':
            return self._docx_to_html(file_path, title)
        if file_ext in {'.xlsx', '.xls'}:
            return self._spreadsheet_to_html(file_path, title)
        return None

    def _page_html(self, title: str, body: str) -> str:
        safe_title = escape(title)
        return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{safe_title}</title>
<style>
  body {{
    font-family: "Segoe UI", "PingFang SC", "Noto Sans SC", "Microsoft YaHei", sans-serif;
    max-width: 920px;
    margin: 32px auto;
    padding: 0 20px 48px;
    color: #1a1a1a;
    line-height: 1.7;
  }}
  h1, h2, h3 {{ line-height: 1.35; }}
  table {{ border-collapse: collapse; width: 100%; margin: 16px 0; }}
  td, th {{ border: 1px solid #d0d0d0; padding: 6px 10px; vertical-align: top; text-align: left; }}
  p {{ margin: 0.6em 0; white-space: pre-wrap; }}
</style>
</head>
<body>
{body}
</body>
</html>"""

    def _docx_to_html(self, file_path: str, title: str) -> str:
        document = docx.Document(file_path)
        blocks: list[str] = []
        for child in document.element.body.iterchildren():
            if child.tag == qn('w:p'):
                paragraph = Paragraph(child, document)
                text = paragraph.text
                if not text.strip():
                    continue
                style_name = ''
                style_id = ''
                if paragraph.style is not None:
                    style_name = (paragraph.style.name or '').lower()
                    style_id = (getattr(paragraph.style, 'style_id', '') or '').lower()
                if style_name in {'title', 'heading 1', '标题 1'} or style_id in {'title', 'heading1'}:
                    tag = 'h1'
                elif style_name in {'heading 2', '标题 2'} or style_id == 'heading2':
                    tag = 'h2'
                elif style_name in {'heading 3', '标题 3'} or style_id == 'heading3':
                    tag = 'h3'
                else:
                    tag = 'p'
                blocks.append(f"<{tag}>{escape(text)}</{tag}>")
            elif child.tag == qn('w:tbl'):
                table = Table(child, document)
                rows = []
                for row in table.rows:
                    cells = ''.join(f"<td>{escape(cell.text)}</td>" for cell in row.cells)
                    rows.append(f"<tr>{cells}</tr>")
                blocks.append("<table>" + "".join(rows) + "</table>")
        body = "\n".join(blocks) or "<p>该文档没有可直接显示的文本内容。</p>"
        return self._page_html(title, body)

    def _spreadsheet_to_html(self, file_path: str, title: str) -> str:
        sheets = pd.read_excel(file_path, sheet_name=None)
        blocks: list[str] = []
        for name, frame in sheets.items():
            blocks.append(f"<h2>{escape(str(name))}</h2>")
            blocks.append(frame.to_html(index=False, na_rep=""))
        body = "\n".join(blocks) or "<p>该表格没有可直接显示的内容。</p>"
        return self._page_html(title, body)
